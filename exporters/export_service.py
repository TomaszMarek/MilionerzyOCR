import csv
import math
from PySide6.QtGui import QTextDocument, QPageSize, QPageLayout
from PySide6.QtCore import QMarginsF
from PySide6.QtPrintSupport import QPrinter

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


class ExportService:
    @staticmethod
    def to_pdf(entries, output_path: str):
        doc = QTextDocument()

        html = """
        <html>
        <head>
            <style>
                body { font-family: 'Segoe UI', Arial, sans-serif; font-size: 10.5pt; color: #1e1e1e; line-height: 1.3; }
                h1 { text-align: center; color: #0b1a30; font-size: 18pt; margin-bottom: 5px; }
                .subtitle { text-align: center; color: #666666; font-size: 10pt; margin-bottom: 25px; }
                .q-block { margin-bottom: 16px; page-break-inside: avoid; }
                .q-title { font-weight: bold; font-size: 11pt; color: #0b1a30; margin-bottom: 6px; }
                .ans-table { width: 100%; border: none; margin-left: 15px; margin-bottom: 4px; }
                .ans-cell { width: 50%; padding: 3px 6px; border: none; font-size: 10pt; }
                .letter { font-weight: bold; color: #333333; }
                hr { border: none; border-top: 1px solid #e0e0e0; margin: 15px 0; }
                .answer-key-header { margin-top: 30px; font-size: 13pt; font-weight: bold; color: #0b1a30; border-bottom: 2px solid #0b1a30; padding-bottom: 4px; }
                .key-grid { width: 100%; margin-top: 10px; border-collapse: collapse; }
                .key-cell { padding: 4px 8px; border: 1px solid #cccccc; text-align: center; font-size: 9.5pt; }
                .key-cell-head { background-color: #f2f2f2; font-weight: bold; }
            </style>
        </head>
        <body>
            <h1>TEST WIEDZY – MILIONERZY</h1>
            <div class="subtitle">Zaznacz prawidłową odpowiedź przy każdym pytaniu</div>
        """

        for idx, e in enumerate(entries, start=1):
            opts = e.warianty or {}
            html += f"""
            <div class="q-block">
                <div class="q-title">{idx}. {e.pytanie}</div>
                <table class="ans-table">
                    <tr>
                        <td class="ans-cell"><span class="letter">[ ] A:</span> {opts.get('A', '')}</td>
                        <td class="ans-cell"><span class="letter">[ ] B:</span> {opts.get('B', '')}</td>
                    </tr>
                    <tr>
                        <td class="ans-cell"><span class="letter">[ ] C:</span> {opts.get('C', '')}</td>
                        <td class="ans-cell"><span class="letter">[ ] D:</span> {opts.get('D', '')}</td>
                    </tr>
                </table>
            </div>
            """

        html += """
            <hr>
            <div class="answer-key-header">KLUCZ ODPOWIEDZI</div>
            <table class="key-grid">
                <tr>
        """
        ITEMS_PER_ROW = 10
        total = len(entries)
        for r_start in range(0, total, ITEMS_PER_ROW):
            chunk = entries[r_start:r_start + ITEMS_PER_ROW]
            html += "<tr>"
            for offset, _ in enumerate(chunk):
                html += f'<td class="key-cell key-cell-head">{r_start + offset + 1}</td>'
            for _ in range(ITEMS_PER_ROW - len(chunk)):
                html += '<td class="key-cell key-cell-head">-</td>'
            html += "</tr><tr>"
            for _, e in enumerate(chunk):
                html += f'<td class="key-cell" style="font-weight: bold; color: #008822;">{e.poprawna_odpowiedz}</td>'
            for _ in range(ITEMS_PER_ROW - len(chunk)):
                html += '<td class="key-cell">-</td>'
            html += "</tr>"

        html += """
            </table>
        </body>
        </html>
        """

        doc.setHtml(html)

        printer = QPrinter(QPrinter.HighResolution)
        printer.setOutputFormat(QPrinter.PdfFormat)
        printer.setOutputFileName(output_path)

        page_layout = QPageLayout(
            QPageSize(QPageSize.A4),
            QPageLayout.Portrait,
            QMarginsF(12.0, 12.0, 12.0, 12.0),
            QPageLayout.Unit.Millimeter
        )
        printer.setPageLayout(page_layout)
        doc.print_(printer)

    @staticmethod
    def to_docx(entries, output_path: str):
        doc = docx.Document()
        sections = doc.sections
        for section in sections:
            section.top_margin = Inches(0.8)
            section.bottom_margin = Inches(0.8)
            section.left_margin = Inches(0.8)
            section.right_margin = Inches(0.8)

        title = doc.add_heading("TEST WIEDZY – MILIONERZY", level=1)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in title.runs:
            run.font.name = "Segoe UI"
            run.font.size = Pt(18)
            run.font.color.rgb = RGBColor(11, 26, 48)
        sub = doc.add_paragraph("Zaznacz prawidłową odpowiedź przy każdym pytaniu.")
        sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sub.runs[0].font.name = "Segoe UI"
        sub.runs[0].font.size = Pt(10)
        sub.runs[0].font.italic = True
        sub.runs[0].font.color.rgb = RGBColor(100, 100, 100)
        doc.add_paragraph()

        for idx, e in enumerate(entries, start=1):
            p_q = doc.add_paragraph()
            p_q.paragraph_format.space_before = Pt(8)
            p_q.paragraph_format.space_after = Pt(2)
            p_q.paragraph_format.keep_with_next = True
            run_q = p_q.add_run(f"{idx}. {e.pytanie}")
            run_q.font.name = "Segoe UI"
            run_q.font.size = Pt(11)
            run_q.font.bold = True
            run_q.font.color.rgb = RGBColor(11, 26, 48)
            opts = e.warianty or {}
            p_ans1 = doc.add_paragraph()
            p_ans1.paragraph_format.left_indent = Inches(0.3)
            p_ans1.paragraph_format.space_after = Pt(1)
            r_a = p_ans1.add_run(f"[  ] A: {opts.get('A', ''):<35}")
            r_b = p_ans1.add_run(f"[  ] B: {opts.get('B', '')}")
            r_a.font.name = "Segoe UI"
            r_b.font.name = "Segoe UI"
            p_ans2 = doc.add_paragraph()
            p_ans2.paragraph_format.left_indent = Inches(0.3)
            p_ans2.paragraph_format.space_after = Pt(6)
            r_c = p_ans2.add_run(f"[  ] C: {opts.get('C', ''):<35}")
            r_d = p_ans2.add_run(f"[  ] D: {opts.get('D', '')}")
            r_c.font.name = "Segoe UI"
            r_d.font.name = "Segoe UI"
        doc.add_page_break()
        h_key = doc.add_heading("KLUCZ POPRAWNYCH ODPOWIEDZI", level=2)
        h_key.runs[0].font.name = "Segoe UI"
        h_key.runs[0].font.color.rgb = RGBColor(11, 26, 48)
        table = doc.add_table(rows=0, cols=10)
        table.style = 'Table Grid'

        total = len(entries)
        for r_start in range(0, total, 10):
            chunk = entries[r_start:r_start + 10]

            row_num = table.add_row()
            row_ans = table.add_row()

            for i in range(10):
                if i < len(chunk):
                    cell_num = row_num.cells[i]
                    cell_num.text = str(r_start + i + 1)
                    cell_num.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                    cell_num.paragraphs[0].runs[0].font.bold = True

                    cell_ans = row_ans.cells[i]
                    cell_ans.text = str(chunk[i].poprawna_odpowiedz)
                    cell_ans.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                    cell_ans.paragraphs[0].runs[0].font.bold = True
                    cell_ans.paragraphs[0].runs[0].font.color.rgb = RGBColor(0, 136, 34)
                else:
                    row_num.cells[i].text = "-"
                    row_ans.cells[i].text = "-"

        doc.save(output_path)

    @staticmethod
    def to_csv(entries, output_path: str):
        with open(output_path, "w", encoding="utf-8-sig", newline="") as f:
            writer = csv.writer(f, delimiter=";", quoting=csv.QUOTE_MINIMAL)
            writer.writerow(
                ["Lp.", "Treść pytania", "Wariant A", "Wariant B", "Wariant C", "Wariant D", "Poprawna odpowiedź"]
            )
            for idx, e in enumerate(entries, start=1):
                opts = e.warianty or {}
                writer.writerow([
                    idx,
                    e.pytanie,
                    opts.get("A", ""),
                    opts.get("B", ""),
                    opts.get("C", ""),
                    opts.get("D", ""),
                    e.poprawna_odpowiedz
                ])

    @staticmethod
    def to_excel(entries, output_path: str):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Pytania"

        headers = ["Lp.", "Treść pytania", "Wariant A", "Wariant B", "Wariant C", "Wariant D", "Poprawna"]
        ws.append(headers)

        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="0B1A30", end_color="0B1A30", fill_type="solid")
        border = Border(
            left=Side(style='thin', color='DDDDDD'),
            right=Side(style='thin', color='DDDDDD'),
            top=Side(style='thin', color='DDDDDD'),
            bottom=Side(style='thin', color='DDDDDD')
        )

        ws.row_dimensions[1].height = 28
        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        QUESTION_COL_WIDTH = 75

        for idx, e in enumerate(entries, start=1):
            opts = e.warianty or {}
            q_text = str(e.pytanie or "")
            row = [
                idx,
                q_text,
                opts.get("A", ""),
                opts.get("B", ""),
                opts.get("C", ""),
                opts.get("D", ""),
                e.poprawna_odpowiedz
            ]
            ws.append(row)

            r_idx = idx + 1
            lines_count = max(1, math.ceil(len(q_text) / (QUESTION_COL_WIDTH - 5)))
            ws.row_dimensions[r_idx].height = max(24, lines_count * 19)

            for c_idx in range(1, len(headers) + 1):
                c = ws.cell(row=r_idx, column=c_idx)
                c.border = border
                c.alignment = Alignment(vertical="center", wrap_text=True)
                if c_idx in (1, 7):
                    c.alignment = Alignment(horizontal="center", vertical="center")
                    if c_idx == 7:
                        c.font = Font(name="Segoe UI", size=11, bold=True, color="008822")

        ws.column_dimensions["A"].width = 6
        ws.column_dimensions["B"].width = QUESTION_COL_WIDTH
        ws.column_dimensions["C"].width = 28
        ws.column_dimensions["D"].width = 28
        ws.column_dimensions["E"].width = 28
        ws.column_dimensions["F"].width = 28
        ws.column_dimensions["G"].width = 14

        wb.save(output_path)

    @staticmethod
    def to_anki_txt(entries, output_path: str):
        with open(output_path, "w", encoding="utf-8") as f:
            for e in entries:
                opts = e.warianty or {}
                front = f"{e.pytanie}<br><br>A: {opts.get('A', '')}<br>B: {opts.get('B', '')}<br>C: {opts.get('C', '')}<br>D: {opts.get('D', '')}"
                correct_letter = e.poprawna_odpowiedz
                correct_text = opts.get(correct_letter, "")
                back = f"<b>{correct_letter}: {correct_text}</b>"
                f.write(f"{front}\t{back}\n")